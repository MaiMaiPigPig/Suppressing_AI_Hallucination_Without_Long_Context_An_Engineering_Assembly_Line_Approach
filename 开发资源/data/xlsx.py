import openpyxl
import csv
import os
import sys

# 配置：修改这里指定你的Excel文件路径
XLSX_PATH = r'f:\跑团工具\开发资源\data\dnd5r.xlsx'

# 自动推导CSV输出目录（同目录下创建csv文件夹）
OUTPUT_DIR = os.path.join(os.path.dirname(XLSX_PATH), 'csv导出1')

def read_and_convert():
    """读取Excel所有工作表并分别转换为CSV，同时打印内容到控制台"""
    if not os.path.exists(XLSX_PATH):
        print(f'错误：找不到文件 {XLSX_PATH}')
        print('请修改脚本中的 XLSX_PATH 变量指向正确的文件路径')
        sys.exit(1)

    # 读取Excel
    wb = openpyxl.load_workbook(XLSX_PATH)
    
    print('=' * 60)
    print(f'Excel文件: {XLSX_PATH}')
    print(f'工作表数量: {len(wb.sheetnames)}')
    print(f'工作表列表: {wb.sheetnames}')
    print('=' * 60)
    print()

    # 创建输出目录
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    print(f'CSV输出目录: {OUTPUT_DIR}')
    print()

    # 遍历每个工作表
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        print('-' * 60)
        print(f'【工作表: {sheet_name}】')
        print(f'  行数: {ws.max_row}, 列数: {ws.max_column}')
        print('-' * 60)
        
        # 打印表头
        if ws.max_row >= 1:
            headers = [cell.value for cell in ws[1]]
            print(f'  表头: {headers}')
        
        # 打印数据（最多显示20行，避免太长）
        data_rows = list(ws.iter_rows(min_row=2, max_row=ws.max_row))
        display_count = min(20, len(data_rows))
        
        if data_rows:
            print(f'  数据（共{len(data_rows)}行，显示前{display_count}行）:')
            for i, row in enumerate(data_rows[:display_count], start=2):
                values = [cell.value for cell in row]
                # 过滤掉全空的行
                if any(v is not None and str(v).strip() != '' for v in values):
                    print(f'    第{i}行: {values}')
            if len(data_rows) > display_count:
                print(f'    ... 还有 {len(data_rows) - display_count} 行未显示')
        else:
            print('  （无数据）')
        print()

        # 保存为CSV（工作表名作为文件名）
        safe_name = "".join(c for c in sheet_name if c.isalnum() or c in (' ', '-', '_')).strip()
        csv_filename = f'{safe_name}.csv'
        csv_path = os.path.join(OUTPUT_DIR, csv_filename)
        
        with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
                writer.writerow(row)
        
        print(f'  ✓ 已保存: {csv_path}')
        print()

    print('=' * 60)
    print('全部导出完成！')
    print('=' * 60)

if __name__ == '__main__':
    read_and_convert()